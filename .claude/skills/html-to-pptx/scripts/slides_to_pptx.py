"""Assembles PNG slide screenshots into a widescreen PPTX (standard 16:9)."""

import sys
import glob
import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu

# Standard PowerPoint widescreen dimensions
SLIDE_W = Emu(12192000)  # 13.333 inches
SLIDE_H = Emu(6858000)   # 7.5 inches


def build_pptx(slides_dir: str, output_path: str) -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    blank_layout = prs.slide_layouts[6]  # completely blank

    pngs = sorted(glob.glob(os.path.join(slides_dir, "slide_*.png")))
    if not pngs:
        sys.exit(f"No slide_*.png files found in {slides_dir}")

    print(f"Packaging {len(pngs)} slides → {output_path}")

    for i, png in enumerate(pngs, 1):
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(png, 0, 0, SLIDE_W, SLIDE_H)
        print(f"  [{i}/{len(pngs)}] {Path(png).name}")

    prs.save(output_path)
    print(f"\nSaved: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python slides_to_pptx.py <slides-dir> <output.pptx>")
        sys.exit(1)
    build_pptx(sys.argv[1], sys.argv[2])
